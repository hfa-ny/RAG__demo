import csv
from pathlib import Path

from langchain_core.documents import Document


SUPPORTED_EXTENSIONS = {".txt", ".md", ".html", ".htm", ".pdf", ".docx", ".csv", ".xlsx", ".pptx"}
FALLBACK_POLICY = (
    "CUNY Demo Policy: All student data must remain on secure, localized servers. "
    "Public LLM APIs are strictly prohibited for processing FERPA-protected information."
)


def _document(text: str, root: Path, file_path: Path, section: str | None = None) -> Document | None:
    content = text.strip()
    if not content:
        return None
    metadata = {
        "source": str(file_path.relative_to(root)),
        "format": file_path.suffix.lower().lstrip("."),
    }
    if section:
        metadata["section"] = section
    return Document(page_content=content, metadata=metadata)


def _load_text(root: Path, file_path: Path) -> list[Document]:
    doc = _document(file_path.read_text(encoding="utf-8"), root, file_path)
    return [doc] if doc else []


def _load_html(root: Path, file_path: Path) -> list[Document]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(file_path.read_text(encoding="utf-8"), "html.parser")
    for element in soup(["script", "style", "noscript"]):
        element.decompose()
    doc = _document(soup.get_text("\n"), root, file_path)
    return [doc] if doc else []


def _load_pdf(root: Path, file_path: Path) -> list[Document]:
    from pypdf import PdfReader

    reader = PdfReader(str(file_path))
    docs: list[Document] = []
    for index, page in enumerate(reader.pages, start=1):
        doc = _document(page.extract_text() or "", root, file_path, f"page {index}")
        if doc:
            docs.append(doc)
    return docs


def _load_docx(root: Path, file_path: Path) -> list[Document]:
    from docx import Document as DocxDocument

    document = DocxDocument(str(file_path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()))
    doc = _document("\n".join(parts), root, file_path, "document")
    return [doc] if doc else []


def _format_rows(rows: list[list[str]], offset: int) -> str:
    lines = []
    for row_index, row in enumerate(rows, start=offset):
        cells = [cell.strip() for cell in row if cell.strip()]
        if cells:
            lines.append(f"Row {row_index}: " + " | ".join(cells))
    return "\n".join(lines)


def _load_csv(root: Path, file_path: Path) -> list[Document]:
    docs: list[Document] = []
    with file_path.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = [[cell for cell in row] for row in csv.reader(handle)]
    for offset in range(0, len(rows), 50):
        batch = rows[offset:offset + 50]
        section = f"rows {offset + 1}-{offset + len(batch)}"
        doc = _document(_format_rows(batch, offset + 1), root, file_path, section)
        if doc:
            docs.append(doc)
    return docs


def _load_xlsx(root: Path, file_path: Path) -> list[Document]:
    from openpyxl import load_workbook

    workbook = load_workbook(file_path, read_only=True, data_only=True)
    docs: list[Document] = []
    for sheet in workbook.worksheets:
        rows = [
            ["" if value is None else str(value) for value in row]
            for row in sheet.iter_rows(values_only=True)
        ]
        for offset in range(0, len(rows), 50):
            batch = rows[offset:offset + 50]
            section = f"{sheet.title} rows {offset + 1}-{offset + len(batch)}"
            doc = _document(_format_rows(batch, offset + 1), root, file_path, section)
            if doc:
                docs.append(doc)
    workbook.close()
    return docs


def _load_pptx(root: Path, file_path: Path) -> list[Document]:
    from pptx import Presentation

    presentation = Presentation(str(file_path))
    docs: list[Document] = []
    for index, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                parts.append(shape.text)
        doc = _document("\n".join(parts), root, file_path, f"slide {index}")
        if doc:
            docs.append(doc)
    return docs


def load_supported_documents(directory: Path) -> list[Document]:
    if not directory.exists():
        directory.mkdir(parents=True)
        (directory / "sample_policy.txt").write_text(FALLBACK_POLICY, encoding="utf-8")

    loaders = {
        ".txt": _load_text,
        ".md": _load_text,
        ".html": _load_html,
        ".htm": _load_html,
        ".pdf": _load_pdf,
        ".docx": _load_docx,
        ".csv": _load_csv,
        ".xlsx": _load_xlsx,
        ".pptx": _load_pptx,
    }
    documents: list[Document] = []
    for file_path in sorted(path for path in directory.rglob("*") if path.is_file()):
        loader = loaders.get(file_path.suffix.lower())
        if not loader:
            continue
        try:
            documents.extend(loader(directory, file_path))
        except Exception as exc:
            relative = file_path.relative_to(directory)
            raise ValueError(f"Unable to extract text from {relative}: {exc}") from exc

    if not documents:
        extensions = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"Add a readable supported document to demo_docs ({extensions}), then try again.")
    return documents
